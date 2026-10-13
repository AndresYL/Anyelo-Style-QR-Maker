"""Lectura de Excel (.xlsx) y CSV para generar QR por filas."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from .core import Item, normalize_data

XLSX_EXT = {".xlsx", ".xlsm"}
CSV_EXT = {".csv", ".tsv", ".txt"}


def sheet_names(path: str | Path) -> list[str]:
    if Path(path).suffix.lower() not in XLSX_EXT:
        return []
    from openpyxl import load_workbook
    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _read_xlsx(p: Path, sheet: str | None) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_csv(p: Path) -> list[list[str]]:
    raw = p.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _clean_headers(row: list[str]) -> list[str]:
    out, seen = [], set()
    for i, h in enumerate(row):
        h = h.strip() or f"Columna {i + 1}"
        base, k = h, 2
        while h in seen:
            h = f"{base} ({k})"
            k += 1
        seen.add(h)
        out.append(h)
    return out


def read_table(path: str | Path, sheet: str | None = None) -> tuple[list[str], list[list[str]]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext in XLSX_EXT:
        rows = _read_xlsx(p, sheet)
    elif ext in CSV_EXT:
        rows = _read_csv(p)
    elif ext == ".xls":
        raise ValueError("Los .xls antiguos no se pueden leer. Abre el archivo y guárdalo como .xlsx o .csv.")
    else:
        raise ValueError("Formato no soportado. Usa .xlsx, .csv o .tsv.")
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ValueError("El archivo no tiene datos.")
    headers = _clean_headers(rows[0])
    w = len(headers)
    return headers, [(r + [""] * w)[:w] for r in rows[1:]]


def guess_columns(headers: list[str], rows: list[list[str]]) -> tuple[int, int]:
    """Adivina (columna_nombre, columna_dato). nombre = -1 significa 'sin nombre'."""
    low = [h.lower() for h in headers]
    data_idx = next((i for i, h in enumerate(low) if any(k in h for k in ("link", "url", "enlace", "qr", "dato"))), None)
    if data_idx is None:
        for i in range(len(headers)):
            col = [r[i] for r in rows if r[i]]
            if col and sum(c.lower().startswith(("http", "www.")) for c in col) / len(col) > 0.5:
                data_idx = i
                break
    if data_idx is None:
        data_idx = len(headers) - 1 if len(headers) > 1 else 0
    name_idx = next((i for i, h in enumerate(low) if i != data_idx and any(k in h for k in ("nombre", "name", "archivo", "título", "titulo"))), None)
    if name_idx is None:
        name_idx = next((i for i in range(len(headers)) if i != data_idx), -1)
    return name_idx, data_idx


def items_from_table(rows: list[list[str]], name_idx: int, data_idx: int, add_https: bool = True) -> list[Item]:
    items = []
    for r in rows:
        data = normalize_data(r[data_idx], add_https)
        if not data:
            continue
        name = r[name_idx].strip() if name_idx >= 0 and r[name_idx].strip() else None
        items.append(Item(name, data))
    return items
